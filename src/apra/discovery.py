"""Safe APRA archive extraction and workbook discovery."""

from __future__ import annotations

import zipfile
from pathlib import Path


def safe_extract(zip_path: Path, destination: Path) -> list[Path]:
    destination = destination.resolve()
    destination.mkdir(parents=True, exist_ok=True)
    extracted: list[Path] = []
    with zipfile.ZipFile(zip_path) as archive:
        for member in archive.infolist():
            target = (destination / member.filename).resolve()
            if destination not in target.parents and target != destination:
                raise ValueError(f"Unsafe ZIP member: {member.filename}")
            if member.is_dir():
                continue
            target.parent.mkdir(parents=True, exist_ok=True)
            with archive.open(member) as source, target.open("wb") as output:
                output.write(source.read())
            extracted.append(target)
    return extracted


def discover_workbooks(source: Path) -> list[Path]:
    if source.is_file() and source.suffix.lower() == ".zip":
        raise ValueError("Extract ZIP archives before workbook discovery")
    return sorted(path for path in source.rglob("*.xlsx") if not path.name.startswith("~$"))


def list_worksheets(workbook: Path) -> list[str]:
    from openpyxl import load_workbook

    book = load_workbook(workbook, read_only=True, data_only=True)
    try:
        return list(book.sheetnames)
    finally:
        book.close()
